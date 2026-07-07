"""
Baixa concluintes 3o ano EM (rede estadual MS) do SGDE e consolida na planilha SED.

API (requer sessao autenticada no navegador):
  GET /api/Sgde/Relatorio/GetRelatorioParaGerencialMatriculaConsolidada
      ?anoFaseId=3&anoReferencia={ano}&cursoId=2&formato=Excel
      &tipoCursoId=117&tipoEnsinoId=1&tipoId=4

Uso:
  # Copie o header Cookie do DevTools (logado em sgde.ms.gov.br/sgde/relatorio)
  set SGDE_COOKIE=.AspNet.ApplicationCookie=...
  python fetch_concluintes_sgde.py

  # Ou informe arquivos Excel ja baixados manualmente:
  python fetch_concluintes_sgde.py --from-dir ../dados/sgde

  # So anos especificos:
  python fetch_concluintes_sgde.py --anos 2013 2014 2015 2016 2017 2018
"""
from __future__ import annotations

try:
    from dotenv import load_dotenv
    load_dotenv()
except ImportError:
    pass

import argparse
import io
import json
import logging
import os
import re
import sys
import time
import urllib.error
import urllib.parse
import urllib.request
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook

sys.path.insert(0, os.path.dirname(__file__))

from enem_config import ANO_FINAL, PASTA_DADOS, configure_logging, resolver_concluintes_xlsx
from enem_helpers import COL_MUNICIPIO, _col_por_prefixo

logger = configure_logging(__name__)

SGDE_BASE = "https://www.sgde.ms.gov.br/api/Sgde/Relatorio/GetRelatorioParaGerencialMatriculaConsolidada"
SGDE_PARAMS = {
    "anoFaseId": 3,
    "cursoId": 2,
    "formato": "Excel",
    "tipoCursoId": 117,
    "tipoEnsinoId": 1,
    "tipoId": 4,
}
ANOS_SGDE_DEFAULT = list(range(2013, 2019))
SHEET_CONSOLIDADA = f"{ANO_FINAL}-2013"
CACHE_DIR = PASTA_DADOS / "sgde"
COLS_CONSOLIDADA = [
    COL_MUNICIPIO,
    "Unidade Escolar",
    "CO_ESCOLA",
    "Localização",
    "Turno",
    "Turma",
    "Concluintes",
    "NU_ANO",
]


def _sgde_cookie() -> str:
    return os.getenv("SGDE_COOKIE", "").strip()


def _sgde_headers(cookie: str) -> dict[str, str]:
    return {
        "accept": "application/json, text/plain, */*",
        "accept-language": "pt-BR,pt;q=0.9",
        "referer": "https://www.sgde.ms.gov.br/sgde/relatorio",
        "user-agent": (
            "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 "
            "(KHTML, like Gecko) Chrome/131.0.0.0 Safari/537.36"
        ),
        "cookie": cookie,
    }


def _is_excel_bytes(data: bytes) -> bool:
    """SGDE pode retornar .xlsx (PK) ou .xls legado (OLE D0CF)."""
    return data[:2] == b"PK" or data[:4] == b"\xd0\xcf\x11\xe0"


def _excel_engine(content: bytes) -> str | None:
    if content[:2] == b"PK":
        return "openpyxl"
    if content[:4] == b"\xd0\xcf\x11\xe0":
        return "xlrd"
    return None


def _read_excel_bytes(content: bytes, **kwargs) -> pd.DataFrame:
    engine = _excel_engine(content)
    if engine is None:
        raise ValueError("Bytes nao parecem Excel (.xls ou .xlsx)")
    return pd.read_excel(io.BytesIO(content), engine=engine, **kwargs)


def baixar_ano_sgde(ano: int, cookie: str, destino: Path | None = None) -> bytes:
    params = {**SGDE_PARAMS, "anoReferencia": ano}
    url = SGDE_BASE + "?" + urllib.parse.urlencode(params)
    req = urllib.request.Request(url, headers=_sgde_headers(cookie), method="GET")
    try:
        with urllib.request.urlopen(req, timeout=120) as resp:
            data = resp.read()
    except urllib.error.HTTPError as exc:
        body = exc.read()[:500]
        raise RuntimeError(f"SGDE HTTP {exc.code} ano {ano}: {body!r}") from exc

    if not _is_excel_bytes(data):
        try:
            msg = json.loads(data.decode("utf-8", errors="replace"))
        except json.JSONDecodeError:
            msg = data[:200]
        raise RuntimeError(f"Resposta inesperada SGDE {ano} (esperado Excel): {msg}")

    if destino:
        destino.parent.mkdir(parents=True, exist_ok=True)
        if destino.suffix.lower() not in (".xls", ".xlsx") and _excel_engine(data) == "xlrd":
            destino = destino.with_suffix(".xls")
        destino.write_bytes(data)
        logger.info("[sgde] %s -> %s (%s KB)", ano, destino, len(data) // 1024)
    return data


def _cell_text(val) -> str:
    if pd.isna(val):
        return ""
    return str(val).strip().lower()


def _detectar_cabecalho(df_preview: pd.DataFrame) -> int:
    for i in range(min(20, len(df_preview))):
        row = df_preview.iloc[i]
        txt = " ".join(_cell_text(x) for x in row.tolist())
        if "munic" in txt and ("matr" in txt or "unidade" in txt):
            return i
    return 0


def parse_sgde_excel(content: bytes, ano: int) -> pd.DataFrame:
    preview = _read_excel_bytes(content, header=None, nrows=20)
    hdr = _detectar_cabecalho(preview)
    df = _read_excel_bytes(content, header=hdr)
    df = df.dropna(how="all")
    df.columns = [str(c).strip() for c in df.columns]

    col_mun = _col_por_prefixo(df.columns, "MUNICIPIO") or COL_MUNICIPIO
    col_esc = _col_por_prefixo(df.columns, "UNIDADE ESCOLAR") or "Unidade Escolar"
    col_loc = _col_por_prefixo(df.columns, "LOCALIZ") or "Localização"
    col_turno = "Turno" if "Turno" in df.columns else _col_por_prefixo(df.columns, "TURNO")
    col_turma = "Turma" if "Turma" in df.columns else _col_por_prefixo(df.columns, "TURMA")
    col_mat = (
        _col_por_prefixo(df.columns, "MATRICULA FINAL")
        or _col_por_prefixo(df.columns, "MATRÍCULA FINAL")
        or _col_por_prefixo(df.columns, "MATRICULA")
    )
    if col_mat is None:
        raise ValueError(f"Coluna de matricula nao encontrada (ano {ano}): {list(df.columns)}")

    out = pd.DataFrame({
        COL_MUNICIPIO: df[col_mun],
        "Unidade Escolar": df[col_esc],
        "Localização": df[col_loc] if col_loc else pd.NA,
        "Turno": df[col_turno] if col_turno else pd.NA,
        "Turma": df[col_turma] if col_turma else pd.NA,
        "Concluintes": pd.to_numeric(df[col_mat], errors="coerce").fillna(0).astype(int),
        "NU_ANO": ano,
        "CO_ESCOLA": pd.NA,
    })
    mun_ok = out[COL_MUNICIPIO].map(_cell_text).str.len() > 0
    mun_ok &= ~out[COL_MUNICIPIO].map(_cell_text).isin({"qtde", "município", "municipio", "munic\u00edpio"})
    esc_ok = out["Unidade Escolar"].map(_cell_text).str.len() > 0
    esc_ok &= ~out["Unidade Escolar"].map(_cell_text).isin({"total", "qtde"})
    out = out[mun_ok & esc_ok & (out["Concluintes"] > 0)].copy()
    if out.empty:
        logger.warning("SGDE %s: nenhuma linha de escola (relatorio vazio ou so totais)", ano)
    return out


def _carregar_base_consolidada(path: Path) -> pd.DataFrame:
    xl = pd.ExcelFile(path)
    for nome in (SHEET_CONSOLIDADA, f"{ANO_FINAL}-2019", "2025-2019"):
        if nome in xl.sheet_names:
            df = pd.read_excel(path, sheet_name=nome)
            df["NU_ANO"] = pd.to_numeric(df["NU_ANO"], errors="coerce")
            return df
    raise ValueError(f"Nenhuma aba consolidada em {path}")


def consolidar_planilha(
    path: Path,
    por_ano: dict[int, pd.DataFrame],
    base_existente: pd.DataFrame | None = None,
) -> None:
    """Grava abas anuais + aba consolidada {ANO_FINAL}-2013."""
    if base_existente is None:
        base_existente = _carregar_base_consolidada(path)

    base = base_existente.copy()
    base = base[~base["NU_ANO"].isin(list(por_ano.keys()))]

    novos = pd.concat([df for df in por_ano.values()], ignore_index=True)
    for col in COLS_CONSOLIDADA:
        if col not in base.columns:
            base[col] = pd.NA
        if col not in novos.columns:
            novos[col] = pd.NA
    consolidado = pd.concat([novos[COLS_CONSOLIDADA], base[COLS_CONSOLIDADA]], ignore_index=True)
    consolidado = consolidado.sort_values([COL_MUNICIPIO, "Unidade Escolar", "NU_ANO", "Turma"])

    wb = load_workbook(path)
    for ano in por_ano:
        nome = str(ano)
        if nome in wb.sheetnames:
            del wb[nome]
    for old in (SHEET_CONSOLIDADA, f"{ANO_FINAL}-2019", "2025-2019"):
        if old in wb.sheetnames and old != SHEET_CONSOLIDADA:
            del wb[old]
    wb.save(path)

    with pd.ExcelWriter(path, engine="openpyxl", mode="a", if_sheet_exists="replace") as writer:
        for ano, df in sorted(por_ano.items()):
            raw_cols = [c for c in df.columns if c != "CO_ESCOLA" or c in df.columns]
            df.to_excel(writer, sheet_name=str(ano), index=False)
        consolidado.to_excel(writer, sheet_name=SHEET_CONSOLIDADA, index=False)

    totais = consolidado.groupby("NU_ANO")["Concluintes"].sum()
    logger.info("[consolidado] %s: %s linhas", SHEET_CONSOLIDADA, len(consolidado))
    logger.info("Totais por ano:\n%s", totais.to_string())


def _parse_args() -> argparse.Namespace:
    p = argparse.ArgumentParser(description="Concluintes SGDE 2013-2018 -> planilha SED")
    p.add_argument("--anos", type=int, nargs="+", default=ANOS_SGDE_DEFAULT)
    p.add_argument("--from-dir", type=Path, help="Pasta com {ano}.xlsx baixados manualmente")
    p.add_argument("--xlsx", type=Path, default=None, help="Planilha destino (default: Concluintes EM ...)")
    p.add_argument("--cookie", default=None, help="Header Cookie SGDE (ou SGDE_COOKIE no .env)")
    p.add_argument("--dry-run", action="store_true", help="Baixa/parseia sem gravar planilha")
    return p.parse_args()


def main() -> None:
    args = _parse_args()
    xlsx = Path(args.xlsx or resolver_concluintes_xlsx())
    if not xlsx.exists():
        raise SystemExit(f"Planilha nao encontrada: {xlsx}")

    por_ano: dict[int, pd.DataFrame] = {}

    if args.from_dir:
        pasta = Path(args.from_dir)
        for ano in args.anos:
            candidatos = [pasta / f"{ano}.xlsx", pasta / f"matricula_{ano}.xlsx", pasta / f"{ano}.xls"]
            path = next((p for p in candidatos if p.exists()), None)
            if path is None:
                logger.warning("Arquivo ausente para %s em %s", ano, pasta)
                continue
            por_ano[ano] = parse_sgde_excel(path.read_bytes(), ano)
            logger.info("%s: %s linhas (arquivo local)", ano, len(por_ano[ano]))
    else:
        cookie = (args.cookie or _sgde_cookie()).strip()
        if not cookie:
            raise SystemExit(
                "Cookie SGDE ausente. Logue em https://www.sgde.ms.gov.br/sgde/relatorio,\n"
                "copie o header Cookie do DevTools > Network e defina SGDE_COOKIE no .env\n"
                "ou use --from-dir com Excel baixados manualmente."
            )
        CACHE_DIR.mkdir(parents=True, exist_ok=True)
        for ano in args.anos:
            cache_xls = CACHE_DIR / f"matricula_consolidada_{ano}.xls"
            cache_xlsx = CACHE_DIR / f"matricula_consolidada_{ano}.xlsx"
            cache = cache_xls if cache_xls.exists() else (cache_xlsx if cache_xlsx.exists() else cache_xls)
            if cache.exists() and cache.stat().st_size > 1000:
                data = cache.read_bytes()
                logger.info("[cache] %s", cache)
            else:
                data = baixar_ano_sgde(ano, cookie, cache)
                time.sleep(0.5)
            por_ano[ano] = parse_sgde_excel(data, ano)
            logger.info("%s: %s linhas, %s concluintes", ano, len(por_ano[ano]), por_ano[ano]["Concluintes"].sum())

    if not por_ano:
        raise SystemExit("Nenhum ano processado.")

    if args.dry_run:
        for ano, df in sorted(por_ano.items()):
            logger.info("DRY %s: %s escolas-turmas, total %s", ano, len(df), df["Concluintes"].sum())
        return

    consolidar_planilha(xlsx, por_ano)
    logger.info("[ok] Planilha atualizada: %s", xlsx)
    logger.info("Proximo passo: python gerar_agregados.py && python gerar_web_data.py")


if __name__ == "__main__":
    main()
